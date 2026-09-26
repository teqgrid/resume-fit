#!/usr/bin/env python3
"""Build the ResumeFit seminar deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "ResumeFit-Seminar.pptx"
LOGO = ROOT / "src" / "assets" / "logo.png"
MUSIC = ROOT / "docs" / "seminar-music.mp3"
GEMINI_DOCS_SHOT = ROOT / "docs" / "gemini-api-key-docs.png"

GEMINI_KEY_URL = "https://aistudio.google.com/apikey"
GEMINI_DOCS_URL = "https://ai.google.dev/gemini-api/docs/api-key"
GEMINI_QUICKSTART_URL = "https://ai.google.dev/gemini-api/docs/quickstart"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
BLUE_DARK = RGBColor(0x1D, 0x4E, 0xD8)
SLATE = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF3, 0xF6, 0xFB)
LIGHT_BLUE = RGBColor(0xEE, 0xF3, 0xFF)
GOLD = RGBColor(0xD9, 0x77, 0x06)
SILVER = RGBColor(0x64, 0x74, 0x8B)
BRONZE = RGBColor(0xB4, 0x53, 0x09)

W = Inches(13.333)
H = Inches(7.5)


def _set_run(run, text, size=20, bold=False, color=NAVY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color)
    return s


def round_rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color)
    return s


def textbox(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h)


def write_box(slide, l, t, w, h, lines, default_size=20):
    box = textbox(slide, l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, str):
            text, size, bold, color = line, default_size, False, NAVY
        else:
            text, size, bold, color = line[0], line[1], line[2], line[3]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, text, size, bold, color)
    return box


def footer(slide, page, total):
    rect(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    box = textbox(slide, Inches(0.4), Inches(7.22), Inches(10), Inches(0.28))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, "ResumeFit  ·  TEQGRID seminar", 11, False, WHITE)
    num = textbox(slide, Inches(11.6), Inches(7.22), Inches(1.4), Inches(0.28))
    p = num.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    _set_run(run, f"{page}  /  {total}", 11, False, WHITE)


def accent_bar(slide):
    rect(slide, 0, 0, Inches(0.12), H, BLUE)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_music(slide, light=False):
    """Speaker icon that plays the seminar bed when clicked."""
    if not MUSIC.exists():
        return
    poster = str(LOGO) if LOGO.exists() else None
    try:
        slide.shapes.add_movie(
            str(MUSIC),
            Inches(12.35),
            Inches(0.18),
            Inches(0.62),
            Inches(0.62),
            poster_frame_image=poster,
            mime_type="audio/mp3",
        )
    except Exception:
        pass
    hint = textbox(slide, Inches(10.7), Inches(0.28), Inches(1.55), Inches(0.4))
    p = hint.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    _set_run(run, "♪  play", 12, True, WHITE if light else BLUE)


def title_slide(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, BLUE)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(0.7), Inches(0.7), Inches(1.05), Inches(1.05))
    write_box(
        s,
        Inches(0.7),
        Inches(2.05),
        Inches(12),
        Inches(4.2),
        [
            ("TEQGRID  ·  HANDS-ON SEMINAR", 16, True, RGBColor(0x93, 0xC5, 0xFD)),
            ("ResumeFit", 54, True, WHITE),
            ("Match a resume to a job with a real AI API", 26, False, RGBColor(0xCB, 0xD5, 0xE1)),
            ("React Native + Google Gemini  ·  live on a phone", 18, False, RGBColor(0x94, 0xA3, 0xB8)),
            ("Click ♪ on this slide for music", 16, False, RGBColor(0x93, 0xC5, 0xFD)),
        ],
    )
    add_music(s, light=True)
    return s


def section_slide(prs, kicker, title, time_label, page, total, music=False):
    s = blank(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, BLUE)
    write_box(
        s,
        Inches(0.8),
        Inches(2.3),
        Inches(11.5),
        Inches(3.2),
        [
            (time_label.upper(), 16, True, RGBColor(0x93, 0xC5, 0xFD)),
            (kicker, 18, False, RGBColor(0x94, 0xA3, 0xB8)),
            (title, 40, True, WHITE),
        ],
    )
    if music:
        add_music(s, light=True)
    footer(s, page, total)


def content_slide(prs, kicker, title, bullets, page, total, note=None, music=False):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    accent_bar(s)
    write_box(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.35), [(kicker.upper(), 13, True, BLUE)])
    write_box(s, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.7), [(title, 30, True, NAVY)])
    lines = [(f"  {b}", 20, False, SLATE) for b in bullets]
    write_box(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(5.2 if not note else 4.4), lines)
    if note:
        round_rect(s, Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.85), LIGHT_BLUE)
        write_box(s, Inches(0.75), Inches(6.28), Inches(11.8), Inches(0.65), [(note, 16, True, BLUE_DARK)])
    if music:
        add_music(s)
    footer(s, page, total)
    return s


def two_col(prs, kicker, title, left_title, left_items, right_title, right_items, page, total):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    accent_bar(s)
    write_box(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.35), [(kicker.upper(), 13, True, BLUE)])
    write_box(s, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.65), [(title, 28, True, NAVY)])
    round_rect(s, Inches(0.5), Inches(1.45), Inches(5.9), Inches(5.4), WHITE)
    round_rect(s, Inches(6.9), Inches(1.45), Inches(5.9), Inches(5.4), WHITE)
    write_box(s, Inches(0.75), Inches(1.6), Inches(5.5), Inches(0.45), [(left_title, 18, True, BLUE)])
    write_box(s, Inches(7.15), Inches(1.6), Inches(5.5), Inches(0.45), [(right_title, 18, True, BLUE)])
    write_box(
        s,
        Inches(0.75),
        Inches(2.15),
        Inches(5.5),
        Inches(4.4),
        [(f"  {x}", 17, False, SLATE) for x in left_items],
    )
    write_box(
        s,
        Inches(7.15),
        Inches(2.15),
        Inches(5.5),
        Inches(4.4),
        [(f"  {x}", 17, False, SLATE) for x in right_items],
    )
    footer(s, page, total)


def cards_slide(prs, kicker, title, cards, page, total, music=False):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    accent_bar(s)
    write_box(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.35), [(kicker.upper(), 13, True, BLUE)])
    write_box(s, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.65), [(title, 28, True, NAVY)])
    n = len(cards)
    gap = Inches(0.28)
    left = Inches(0.5)
    usable = W - Inches(1.0)
    cw = int((usable - gap * (n - 1)) / n)
    for i, (ct, cb) in enumerate(cards):
        x = left + i * (cw + gap)
        round_rect(s, x, Inches(1.5), cw, Inches(5.35), WHITE)
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.28), Inches(1.75), Inches(0.42), Inches(0.42))
        fill(badge, BLUE)
        tb = textbox(s, x + Inches(0.28), Inches(1.78), Inches(0.42), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        _set_run(run, str(i + 1), 14, True, WHITE)
        write_box(
            s,
            x + Inches(0.25),
            Inches(2.35),
            cw - Inches(0.5),
            Inches(4.2),
            [(ct, 20, True, NAVY), (cb, 16, False, SLATE)],
        )
    if music:
        add_music(s)
    footer(s, page, total)


def add_link(slide, l, t, w, h, label, url):
    """Clickable URL chip — opens the official page from PowerPoint."""
    shape = round_rect(slide, l, t, w, h, LIGHT_BLUE)
    try:
        shape.click_action.hyperlink.address = url
    except Exception:
        pass
    box = textbox(slide, l + Inches(0.1), t + Inches(0.08), w - Inches(0.2), h - Inches(0.1))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    _set_run(run, label, 12, True, BLUE_DARK)
    try:
        run.hyperlink.address = url
    except Exception:
        pass
    return shape


def flow_box(slide, l, t, w, h, title, subtitle, fill_color=WHITE, title_color=NAVY):
    round_rect(slide, l, t, w, h, fill_color)
    write_box(
        slide,
        l + Inches(0.1),
        t + Inches(0.08),
        w - Inches(0.2),
        h - Inches(0.14),
        [(title, 13, True, title_color), (subtitle, 11, False, SLATE)],
    )


def arrow_right(slide, l, t):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, Inches(0.26), Inches(0.16))
    fill(shape, BLUE)
    return shape


def arrow_down(slide, l, t):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, l, t, Inches(0.16), Inches(0.24))
    fill(shape, BLUE)
    return shape


def arrow_left(slide, l, t):
    shape = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, l, t, Inches(0.26), Inches(0.16))
    fill(shape, BLUE)
    return shape


def rn_cycle_slide(prs, page, total):
    """Boot + analyze cycle that matches the current ResumeFit files."""
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    accent_bar(s)
    write_box(s, Inches(0.55), Inches(0.22), Inches(12), Inches(0.3), [("04  ·  THIS CODE", 13, True, BLUE)])
    write_box(
        s,
        Inches(0.55),
        Inches(0.48),
        Inches(12.2),
        Inches(0.5),
        [("React Native app cycle — ResumeFit as it is wired today", 24, True, NAVY)],
    )

    boot = [
        ("1  Native boot", "iOS / Android looks up app.json name ResumeFit"),
        ("2  index.js", "AppRegistry.registerComponent → App"),
        ("3  App.tsx", "GestureHandler + SafeArea + navigator"),
        ("4  navigation", "Stack: Home first, then Result"),
    ]
    analyze = [
        ("5  HomeScreen", "onPickResume + paste JD + onAnalyze"),
        ("6  pick + base64", "pickResume.ts then uriToBase64(pdf)"),
        ("7  gemini.ts", "POST generateContent  ·  Flash Lite JSON"),
        ("8  ResultScreen", "Draws the JSON only. Does not re-score."),
    ]
    left = Inches(0.45)
    bw = Inches(2.85)
    gap = Inches(0.38)
    y1 = Inches(1.15)
    y2 = Inches(3.55)
    for i, (title, sub) in enumerate(boot):
        x = left + i * (bw + gap)
        flow_box(s, x, y1, bw, Inches(1.55), title, sub, WHITE, BLUE)
        if i < 3:
            arrow_right(s, x + bw + Inches(0.06), y1 + Inches(0.68))
    arrow_down(s, left + 3 * (bw + gap) + bw / 2 - Inches(0.08), Inches(2.82))
    for i, (title, sub) in enumerate(analyze):
        x = left + (3 - i) * (bw + gap)
        fill_c = LIGHT_BLUE if i == 2 else WHITE
        flow_box(s, x, y2, bw, Inches(1.55), title, sub, fill_c, BLUE if i == 2 else NAVY)
        if i < 3:
            dest_x = left + (2 - i) * (bw + gap)
            arrow_left(s, dest_x + bw + Inches(0.06), y2 + Inches(0.68))
    write_box(
        s,
        Inches(0.45),
        Inches(5.42),
        Inches(12.4),
        Inches(1.55),
        [
            ("Live path in this repo", 14, True, NAVY),
            (
                "HomeScreen.onAnalyze → uriToBase64 + analyzeResumeMatch (src/api/gemini.ts) → navigation.navigate('Result', { analysis }).",
                14,
                False,
                SLATE,
            ),
            (
                "Key files: index.js · App.tsx · src/navigation/index.tsx · HomeScreen.tsx · pickResume.ts · gemini.ts · ResultScreen.tsx · .env",
                13,
                False,
                SLATE,
            ),
        ],
    )
    footer(s, page, total)


def gemini_key_diagram_slide(prs, page, total):
    """Official Gemini key steps + docs screenshot + clickable URLs."""
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    accent_bar(s)
    write_box(s, Inches(0.55), Inches(0.2), Inches(12), Inches(0.28), [("05  ·  GOOGLE AI STUDIO", 13, True, BLUE)])
    write_box(
        s,
        Inches(0.55),
        Inches(0.44),
        Inches(12.2),
        Inches(0.42),
        [("Gemini key — official steps (click the blue chips)", 22, True, NAVY)],
    )

    steps = [
        ("1", "Sign in with a Google account"),
        ("2", "Open aistudio.google.com/apikey"),
        ("3", "Create API key  ·  pick / create a Cloud project"),
        ("4", "Copy the key once  ·  treat it like a password"),
        ("5", "cp .env.example .env   then   GEMINI_API_KEY=…"),
        ("6", "Restart Metro  ·  never commit .env"),
    ]
    y = Inches(0.95)
    for num, text in steps:
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), y, Inches(0.36), Inches(0.36))
        fill(badge, BLUE)
        tb = textbox(s, Inches(0.5), y + Inches(0.02), Inches(0.36), Inches(0.32))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        _set_run(run, num, 13, True, WHITE)
        round_rect(s, Inches(0.98), y - Inches(0.04), Inches(6.15), Inches(0.46), WHITE)
        write_box(s, Inches(1.12), y, Inches(5.9), Inches(0.38), [(text, 14, False, SLATE)])
        if num != "6":
            arrow_down(s, Inches(0.6), y + Inches(0.36))
        y += Inches(0.58)

    if GEMINI_DOCS_SHOT.exists():
        s.shapes.add_picture(str(GEMINI_DOCS_SHOT), Inches(7.35), Inches(0.95), Inches(5.5), Inches(3.85))
        write_box(
            s,
            Inches(7.35),
            Inches(4.82),
            Inches(5.5),
            Inches(0.35),
            [("Official docs screenshot  ·  ai.google.dev", 11, False, SLATE)],
        )
    else:
        round_rect(s, Inches(7.35), Inches(0.95), Inches(5.5), Inches(3.85), WHITE)
        write_box(
            s,
            Inches(7.55),
            Inches(2.3),
            Inches(5.1),
            Inches(1.2),
            [("Open the docs URL below for the live Create API key page.", 16, False, SLATE)],
        )

    add_link(s, Inches(7.35), Inches(5.2), Inches(5.5), Inches(0.42), "Create key  →  aistudio.google.com/apikey", GEMINI_KEY_URL)
    add_link(s, Inches(7.35), Inches(5.7), Inches(2.65), Inches(0.42), "API key docs", GEMINI_DOCS_URL)
    add_link(s, Inches(10.2), Inches(5.7), Inches(2.65), Inches(0.42), "Quickstart", GEMINI_QUICKSTART_URL)

    write_box(
        s,
        Inches(0.5),
        Inches(6.22),
        Inches(12.3),
        Inches(0.8),
        [
            (
                "Never paste the key in Slack, GitHub, or this deck. Classroom demo keeps it in .env on the phone. Play Store → put the key on YOUR server.",
                13,
                False,
                SLATE,
            ),
        ],
    )
    footer(s, page, total)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 24

    title_slide(prs)

    # 2 agenda
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    accent_bar(s)
    write_box(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.35), [("TODAY", 13, True, BLUE)])
    write_box(s, Inches(0.55), Inches(0.58), Inches(12), Inches(0.6), [("How we will run this room", 28, True, NAVY)])
    rows = [
        ("01", "My intro", "Who is teaching today"),
        ("02", "TEQGRID intro", "Why we are here"),
        ("03", "Project intro", "What ResumeFit does"),
        ("04", "React Native structure", "How a phone app is wired"),
        ("05", "Google API key", "Create your own Gemini key"),
        ("06", "How the app works", "PDF + JD → score"),
        ("07", "Student Q&A", "Your questions"),
        ("08", "10-question quiz", "Top 3 win"),
        ("09", "Thank you", "♪ music on title, quiz, winners"),
    ]
    y = Inches(1.28)
    for t, name, why in rows:
        round_rect(s, Inches(0.5), y, Inches(12.3), Inches(0.58), WHITE)
        write_box(s, Inches(0.7), y + Inches(0.1), Inches(1.1), Inches(0.4), [(t, 15, True, BLUE)])
        write_box(s, Inches(2.0), y + Inches(0.1), Inches(4.6), Inches(0.4), [(name, 16, True, NAVY)])
        write_box(s, Inches(6.8), y + Inches(0.1), Inches(5.7), Inches(0.4), [(why, 15, False, SLATE)])
        y += Inches(0.62)
    add_music(s)
    footer(s, 2, total)

    # 3 my intro
    two_col(
        prs,
        "01  ·  speaker",
        "Hi — I am Aishwarya Rastogi",
        "Who I am",
        [
            "Aishwarya Rastogi",
            "React Native developer",
            "I ship small phone apps with real APIs",
            "No mock screens for this seminar",
            "",
            "Today I will teach from a live app",
            "on a real iPhone — ResumeFit.",
        ],
        "How I will teach",
        [
            "One idea → one file → one question",
            "We will not open 20 files",
            "You will see: input, API, result",
            "",
            "Ask anytime. Wrong answers are useful.",
            "The quiz at the end uses this talk.",
        ],
        3,
        total,
    )

    # 4 teqgrid
    content_slide(
        prs,
        "02  ·  TEQGRID",
        "Why TEQGRID brought you here",
        [
            "• TEQGRID trains builders — not slide readers.",
            "• Line we use in production: frontend → backend → AI.",
            "• Today you learn the AI middle: a phone calls Gemini and shows JSON.",
            "• A backend would hide the key. We skip it so you can see the full feature.",
            "• Goal: leave able to clone, paste your key, and change one prompt line.",
            "• Participation is the point. The quiz is how we pick the top 3.",
        ],
        4,
        total,
        "If TEQGRID is issuing certificates, mention that after the winners slide.",
    )

    # 5 project intro
    two_col(
        prs,
        "03  ·  project",
        "ResumeFit — a junior recruiter on a phone",
        "The product",
        [
            "Upload a PDF resume",
            "Paste a job description",
            "Tap Analyze match",
            "",
            "Gemini reads both",
            "You get a score you can argue with",
        ],
        "What comes back",
        [
            "Match score 0–100",
            "Matching skills (green chips)",
            "Missing skills (red chips)",
            "Experience alignment",
            "3 concrete resume edits",
            "",
            "React Native CLI — not Expo",
        ],
        5,
        total,
    )

    content_slide(
        prs,
        "03  ·  project",
        "This is a feature, not ChatGPT in a browser",
        [
            "• User gives input → our app calls an API → user sees a usable result.",
            "• We are not training a model. We are calling one.",
            "• Weather API returns temperature. Gemini returns text. We force JSON.",
            "• Two inputs only: resume PDF + job description.",
            "• No fake analysis. If the key or the model fails, you see the real error.",
        ],
        6,
        total,
        "Ask: what two inputs does this feature need? Wait for: PDF + JD.",
    )

    # 7-8 RN structure
    section_slide(prs, "04", "How a React Native app is structured", "Application working structure", 7, total)

    cards_slide(
        prs,
        "04  ·  React Native",
        "One JavaScript/TypeScript app → two phones",
        [
            ("JS / TS layer", "Screens, state, fetch(). HomeScreen and ResultScreen live here. Same code for iOS and Android."),
            ("Native layer", "iOS (Swift + Xcode) and Android (Kotlin + Gradle) host the JS. Camera, files, splash, icon."),
            ("Metro", "Dev server on your Mac. The phone loads index.js from 8081. Change JS → Fast Refresh."),
        ],
        8,
        total,
    )

    content_slide(
        prs,
        "04  ·  React Native",
        "Folders you will actually open today",
        [
            "• App.tsx + src/navigation — which screen is showing",
            "• src/screens/HomeScreen.tsx — pick PDF, paste JD, tap Analyze",
            "• src/screens/ResultScreen.tsx — dumb UI. It only draws JSON.",
            "• src/api/pickResume.ts — system file picker, PDF only",
            "• src/api/gemini.ts — URL, key, prompt, fetch, parse",
            "• .env — GEMINI_API_KEY. Never commit this file.",
            "• ios/ and android/ — splash, icons, signing. Rebuild after native changes.",
        ],
        9,
        total,
        "Stay in 3 files for the live walk: pickResume.ts, gemini.ts, HomeScreen.tsx.",
    )

    rn_cycle_slide(prs, 10, total)

    # 11-13 API key
    section_slide(prs, "05", "Generate a Google Gemini API key", "Free key  ·  4 minutes", 11, total)

    content_slide(
        prs,
        "05  ·  Google AI Studio",
        "Create your own key — do not share one room key",
        [
            "1. Open https://aistudio.google.com/apikey  (Google account).",
            "2. Click Create API key. Pick or create a Google Cloud project if asked.",
            "3. Copy the key once. Treat it like a password.",
            "4. In the project:  cp .env.example .env",
            "5. Paste:  GEMINI_API_KEY=your_key   (no quotes, no extra space).",
            "6. Restart Metro. The app reads .env only at start.",
            "7. If 40 people hammer one key you get 429 / quota. Each student = one key.",
        ],
        12,
        total,
        "This classroom key lives on the phone. Play Store? Put the key on YOUR server.",
    )

    gemini_key_diagram_slide(prs, 13, total)

    # 14-17 how app works
    section_slide(prs, "06", "How ResumeFit actually works", "The live path", 14, total)

    content_slide(
        prs,
        "06  ·  flow",
        "What happens when they tap Analyze",
        [
            "1. Phone has a PDF + a job description",
            "2. App converts the PDF to base64 (JSON cannot carry a raw file)",
            "3. POST to Gemini generateContent  (gemini-3.1-flash-lite)",
            "4. Prompt says: recruiter, be honest, do not invent skills, return this JSON",
            "5. Gemini reads PDF + JD and returns score, chips, 3 edits",
            "6. Result screen only displays that JSON — it does not re-score",
        ],
        15,
        total,
        "There is no backend today. The phone talks to Google.",
    )

    two_col(
        prs,
        "06  ·  three files",
        "The working path — do not rewrite the app",
        "The letter",
        [
            "Address = URL (generateContent)",
            "Stamp / ID = API key",
            "Body = prompt + PDF",
            "Reply = JSON",
            "",
            "text = instructions + job",
            "inline_data = PDF as base64",
        ],
        "The files",
        [
            "pickResume.ts — picker, cancel = null",
            "gemini.ts — key, parts, fetch, parse",
            "HomeScreen.tsx — validate, load, navigate",
            "",
            "setLoading(true) before the network.",
            "finally { setLoading(false) } always.",
        ],
        16,
        total,
    )

    content_slide(
        prs,
        "06  ·  prompt",
        "The model follows the prompt, not a magic HR API",
        [
            "• Role: you are a recruiter.",
            "• Task: compare this resume PDF to this job description.",
            "• Constraints: be honest; if not a resume, score 0; do not invent skills.",
            "• Output contract: exact JSON keys the Result screen already uses.",
            "• temperature 0.2 — scoring wants boring and repeatable.",
            "• If green chips show skills that are not on the PDF, fix the prompt, not the CSS.",
        ],
        17,
        total,
        "Live tweak: change “recruiter” to “strict hiring manager.” Same API, different score.",
    )

    # 18 Q&A
    section_slide(prs, "07", "Student Q&A", "Your turn", 18, total, music=True)

    content_slide(
        prs,
        "07  ·  Q&A",
        "Ask anything. If the room is quiet, start here.",
        [
            "• Why PDF, not Word?  —  one path; Gemini document understanding likes PDF.",
            "• Why base64, not Drive?  —  Drive needs login. Base64 = one POST.",
            "• Shared key for 40 students?  —  quota dies. Each person uses their own key.",
            "• Is the score “true”?  —  no. Model opinion under our prompt.",
            "• Match 3 jobs?  —  loop the same function three times.",
            "• Scanned photo resume?  —  often yes. Prefer a real text PDF for the demo.",
        ],
        19,
        total,
        "Write one question on a slip if you do not want to speak.",
        music=True,
    )

    # 20 quiz intro
    section_slide(prs, "08", "10-question test", "Write A / B / C  ·  no phones", 20, total, music=True)

    content_slide(
        prs,
        "08  ·  quiz  1–5",
        "Circle one letter. 30 seconds each.",
        [
            "1. ResumeFit needs which two inputs?   A Photo + recipes   B Resume PDF + job description   C Email + OTP",
            "2. The Gemini key should live in?   A GitHub   B .env (not committed)   C The Result screen",
            "3. Gemini must return what the UI can draw as?   A A video   B JSON   C An Excel file",
            "4. We send the PDF as?   A Base64 inline_data   B A Drive link   C A screenshot",
            "5. This classroom demo has?   A No backend   B A Node server   C A database",
        ],
        21,
        total,
        music=True,
    )

    content_slide(
        prs,
        "08  ·  quiz  6–10",
        "Same rules. Last five.",
        [
            "6. Which file calls Gemini?   A App.tsx   B src/api/gemini.ts   C Info.plist",
            "7. React Native lets you?   A One JS/TS app → iOS + Android   B Websites only   C iOS only",
            "8. Delete “do not invent skills” and green chips may?   A Vanish   B Show skills not on the PDF   C Crash Metro",
            "9. Free key is created at?   A aistudio.google.com/apikey   B github.com   C npmjs.com",
            "10. For Play Store you should?   A Keep the key in the APK   B Put the key on your server   C Print it on splash",
        ],
        22,
        total,
        "Host key: 1B  2B  3B  4A  5A  6B  7A  8B  9A  10B",
        music=True,
    )

    # 23 winners
    s = blank(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, BLUE)
    write_box(
        s,
        Inches(0.7),
        Inches(0.45),
        Inches(12),
        Inches(1.2),
        [
            ("QUIZ WINNERS", 16, True, RGBColor(0x93, 0xC5, 0xFD)),
            ("Top 3 students take the quiz", 34, True, WHITE),
        ],
    )
    medals = [
        ("1st", "Gold", GOLD, "Highest score  ·  walks through one prompt line"),
        ("2nd", "Silver", SILVER, "Next highest  ·  explains URL + key + JSON"),
        ("3rd", "Bronze", BRONZE, "Third  ·  names the 3 files"),
    ]
    for i, (place, metal, color, prize) in enumerate(medals):
        x = Inches(0.7) + i * Inches(4.1)
        round_rect(s, x, Inches(2.1), Inches(3.85), Inches(3.7), WHITE)
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(1.45), Inches(2.35), Inches(0.95), Inches(0.95))
        fill(badge, color)
        tb = textbox(s, x + Inches(1.45), Inches(2.52), Inches(0.95), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        _set_run(run, place, 16, True, WHITE)
        write_box(
            s,
            x + Inches(0.25),
            Inches(3.5),
            Inches(3.35),
            Inches(2.0),
            [
                (metal, 22, True, NAVY),
                ("Name: ____________________", 16, False, SLATE),
                (prize, 14, False, SLATE),
            ],
        )
    add_music(s, light=True)
    footer(s, 23, total)

    # 24 thank you
    s = blank(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, BLUE)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(0.75), Inches(1.15), Inches(0.95), Inches(0.95))
    write_box(
        s,
        Inches(0.75),
        Inches(2.3),
        Inches(12),
        Inches(4.4),
        [
            ("THANK YOU", 16, True, RGBColor(0x93, 0xC5, 0xFD)),
            ("Clone. Paste your key. Change one prompt line.", 30, True, WHITE),
            ("aistudio.google.com/apikey", 22, False, RGBColor(0x93, 0xC5, 0xFD)),
            ("github.com/Aishwaryaofficial/resume-fit", 18, False, RGBColor(0xCB, 0xD5, 0xE1)),
            ("npm start    then    npm run ios  /  npm run android", 18, False, RGBColor(0x94, 0xA3, 0xB8)),
            ("♪ click the note  ·  questions anytime  ·  TEQGRID", 16, False, RGBColor(0x93, 0xC5, 0xFD)),
        ],
    )
    add_music(s, light=True)
    footer(s, 24, total)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
